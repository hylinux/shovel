#---------------------------------------------------------
# Excel 解析器 (openpyxl)
#
# 日期: 2026-09-29
# 作者: HongWei Guo <hongweig@163.com>
#---------------------------------------------------------
"""把工作表切成"带表头的行组"。

表格是所有格式里最容易做错的一种, 错法有两个极端:

* **整表一块** —— 一张三千行的表变成一个巨大的 chunk, embedding 只能
  得到一个模糊的平均语义, 搜"张三的报销"会命中整张表, 然后 LLM 要在
  三千行里自己找。
* **一行一块** —— 每块只有 "张三 | 2024-03 | 1200", 脱离表头之后连
  "1200 是什么"都不知道, 而且三千个碎块会把召回名额全部占满。

正确的中间态是**行组 + 表头重复**: 每 N 行一块, 每块都带上表头行。
代价是表头被存了很多遍 (几十字节), 收益是每一块都能独立理解。

另外两个具体决定:

* ``read_only=True`` + ``values_only=True``: 一个 50 MB 的 xlsx 用普通
  模式打开会在内存里建出全部 Cell 对象, 轻易上 GB。只读模式是流式的。
* ``data_only=True``: 取公式的**计算结果**而不是公式本身。用户搜的是
  "总额 12800", 不是 "=SUM(B2:B30)"。代价是文件若从未被 Excel 打开过
  保存过, 缓存值可能不存在 —— 这种情况下该单元格为空, 属于可接受。
"""

from __future__ import annotations

import datetime as dt
import io
from typing import Any, ClassVar

from shovel.domain.enums import BlockKind, DocClass, Modality
from shovel.exceptions.pipeline import EmptyDocumentError, ParseError
from shovel.pipeline.parsers.base import (
    Block,
    ParsedDocument,
    ParseHint,
    ParserSpec,
    clean_text,
    take,
)

#: 每块包含多少数据行 (不含重复的表头)。
#:
#: 20 行是按"一块落在 embedding 模型舒适区 (几百 token)"倒推的:
#: 典型业务表一行 5~8 列、每格十来个字符, 20 行约 800~1500 字符。
_ROWS_PER_BLOCK = 20

#: 单个工作表最多读多少行。超出就截断 —— 一张十万行的数据表本质是
#: 数据库导出, 全量索引它既慢又对检索没有帮助 (用户真要查会去查库)。
_MAX_ROWS_PER_SHEET = 5000

#: 一个单元格最多保留多少字符。超长单元格通常是整段备注被塞进了格子。
_MAX_CELL_CHARS = 500


class ExcelParser:

    spec: ClassVar[ParserSpec] = ParserSpec(
        name="excel",
        version="1",
        doc_classes=frozenset({DocClass.TABLE}),
        suffixes=frozenset({".xlsx", ".xlsm"}),
    )

    def parse(self, payload: bytes, hint: ParseHint) -> ParsedDocument:
        import openpyxl

        try:
            book = openpyxl.load_workbook(
                io.BytesIO(payload),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise ParseError(
                hint.uri,
                f"无法打开 Excel 工作簿: {exc}",
                hint="确认它是 .xlsx/.xlsm (而非旧版 .xls), 且文件未损坏。",
            ) from exc

        blocks: list[Block] = []
        truncated_sheets: list[str] = []
        # read_only 模式下 close() 会释放底层 zip 句柄, 之后再访问
        # book.worksheets 会抛异常。所以表名必须在关闭前先取出来。
        sheet_names = [str(s.title) for s in book.worksheets]

        try:
            for sheet in book.worksheets:
                # 隐藏表通常是中间计算过程或历史备份, 索引它们只会
                # 让用户搜到自己早就"删掉"的旧数据。
                if sheet.sheet_state != "visible":
                    continue

                sheet_blocks, cut = _sheet_blocks(sheet)
                blocks.extend(sheet_blocks)

                if cut:
                    truncated_sheets.append(str(sheet.title))
        finally:
            book.close()

        if not blocks:
            raise EmptyDocumentError(hint.uri)

        kept, truncated = take(tuple(blocks), hint.max_chars)

        return ParsedDocument(
            blocks=kept,
            title=hint.filename,
            doc_class=DocClass.TABLE,
            parser=self.spec.name,
            parser_version=self.spec.version,
            truncated=truncated or bool(truncated_sheets),
            meta={
                "sheets": sheet_names,
                "truncated_sheets": truncated_sheets,
                "modality": Modality.TABLE.value,
            },
        )


def _sheet_blocks(sheet: Any) -> tuple[list[Block], bool]:
    """一个工作表 -> 若干行组块。返回 (块, 是否因行数上限被截断)。"""

    rows: list[list[str]] = []
    cut = False

    for offset, values in enumerate(sheet.iter_rows(values_only=True)):
        if offset >= _MAX_ROWS_PER_SHEET:
            cut = True
            break

        cells = [_render(v) for v in values]

        # 全空行是表与表之间的视觉分隔, 不是数据
        if any(cells):
            rows.append(cells)

    if not rows:
        return [], cut

    width = max(len(r) for r in rows)
    header = _pad(rows[0], width)

    # 表头判定: 第一行**全部**是非数字才算表头。
    # 只要有一格是数字, 它就更可能是数据行 —— 这时宁可用 "列1/列2"
    # 当表头, 也不能把一行真实数据吃掉当标题。
    has_header = all(
        cell and not _looks_numeric(cell) for cell in header
    )

    if not has_header:
        header = [f"列{i + 1}" for i in range(width)]
        body = rows
    else:
        body = rows[1:]

    if not body:
        # 只有表头没有数据: 仍然值得索引一块, 列名本身是有检索价值的
        body = []

    title = str(sheet.title)
    blocks: list[Block] = []

    rule = "| " + " | ".join(["---"] * width) + " |"
    header_line = "| " + " | ".join(header) + " |"

    if not body:
        return [Block(
            text="\n".join([header_line, rule]),
            kind=BlockKind.TABLE,
            path=(title,),
            locator=f"{title}!1",
            meta={"sheet": title, "cols": width},
        )], cut

    first_data_row = 2 if has_header else 1

    for start in range(0, len(body), _ROWS_PER_BLOCK):
        group = body[start:start + _ROWS_PER_BLOCK]
        lines = [header_line, rule]

        for row in group:
            lines.append("| " + " | ".join(_pad(row, width)) + " |")

        row_from = first_data_row + start
        row_to = row_from + len(group) - 1

        blocks.append(Block(
            text="\n".join(lines),
            kind=BlockKind.TABLE,
            path=(title,),
            # Excel 用户认得 "Sheet1!2:21" 这种写法, 它能直接定位
            locator=f"{title}!{row_from}:{row_to}",
            meta={"sheet": title, "rows": len(group), "cols": width},
        ))

    return blocks, cut


def _render(value: object) -> str:
    """把单元格的值变成字符串。

    日期必须自己格式化: 默认的 ``str(datetime)`` 会给出
    ``2024-03-01 00:00:00``, 那串永远为零的时间部分在每个日期格里
    重复出现, 纯属噪声。
    """

    if value is None:
        return ""

    if isinstance(value, dt.datetime):
        if value.hour or value.minute or value.second:
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return value.strftime("%Y-%m-%d")

    if isinstance(value, dt.date):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, float) and value.is_integer():
        # 1200.0 -> 1200: Excel 里整数也是 float, 保留 .0 会让
        # 用户搜 "1200" 时对不上
        return str(int(value))

    text = clean_text(str(value)).replace("\n", " ").replace("|", "\\|")

    return text[:_MAX_CELL_CHARS]


def _looks_numeric(cell: str) -> bool:
    try:
        float(cell.replace(",", ""))
    except ValueError:
        return False
    return True


def _pad(row: list[str], width: int) -> list[str]:
    return row + [""] * (width - len(row))


PARSER = ExcelParser()

__all__ = ["PARSER", "ExcelParser"]
