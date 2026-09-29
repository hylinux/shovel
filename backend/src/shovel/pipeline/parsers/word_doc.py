#---------------------------------------------------------
# Word 解析器 (python-docx)
#
# 日期: 2026-09-29
# 作者: HongWei Guo <hongweig@163.com>
#---------------------------------------------------------
"""抽取 .docx 的正文、标题层级与表格。

python-docx 的 ``document.paragraphs`` 和 ``document.tables`` 是两个互不
相干的列表, 直接用它们会**丢掉顺序** —— 所有表格会被堆到文末, 于是
"下表列出了三种方案"这句话和它下面那张表就彻底分家了。

正确做法是按 ``document.element.body`` 的子元素顺序走一遍, 遇到 ``w:p``
当段落、遇到 ``w:tbl`` 当表格。多写二十行, 换回正文与表格的相对位置。

标题层级取自段落样式名 (``Heading 1``/``标题 1``), 而不是字号。字号
在实际文档里极不可靠: 很多人用加粗的正文冒充标题, 也有人把标题调小。
样式名虽然也可能缺失, 但它错的时候是"没有", 不是"错的"。
"""

from __future__ import annotations

import io
import re
from typing import Any, ClassVar

from shovel.domain.enums import BlockKind, DocClass
from shovel.exceptions.pipeline import EmptyDocumentError, ParseError
from shovel.pipeline.parsers.base import (
    Block,
    HeadingStack,
    ParsedDocument,
    ParseHint,
    ParserSpec,
    clean_text,
    take,
)

#: 样式名 -> 标题层级。中英文都要认: 中文版 Word 存的就是 "标题 1"。
_HEADING_STYLE = re.compile(r"^(?:heading|标题)\s*([1-9])$", re.IGNORECASE)

#: 列表样式名的特征。用于把列表项标出来, 好让切块阶段把零散的
#: 三五个字的条目粘成一块。
_LIST_STYLE = re.compile(r"list|bullet|number|列表|项目符号", re.IGNORECASE)

_WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class WordParser:

    spec: ClassVar[ParserSpec] = ParserSpec(
        name="word",
        version="1",
        doc_classes=frozenset({DocClass.OFFICE_DOC}),
        suffixes=frozenset({".docx"}),
    )

    def parse(self, payload: bytes, hint: ParseHint) -> ParsedDocument:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        try:
            document = docx.Document(io.BytesIO(payload))
        except Exception as exc:
            # python-docx 对着一个 .doc 会抛 PackageNotFoundError,
            # 对着损坏的 zip 会抛别的。对用户来说都是同一件事。
            raise ParseError(
                hint.uri,
                f"无法打开 Word 文档: {exc}",
                hint="确认它是 .docx (而非旧版 .doc), 且文件未损坏。",
            ) from exc

        stack = HeadingStack()
        blocks: list[Block] = []
        index = 0

        for child in document.element.body.iterchildren():
            tag = child.tag

            if tag == f"{_WORD_NS}p":
                index += 1
                block = _paragraph_block(Paragraph(child, document), stack, index)

                if block is not None:
                    blocks.append(block)

            elif tag == f"{_WORD_NS}tbl":
                index += 1
                block = _table_block(Table(child, document), stack, index)

                if block is not None:
                    blocks.append(block)

        if not blocks:
            raise EmptyDocumentError(hint.uri)

        kept, truncated = take(tuple(blocks), hint.max_chars)

        return ParsedDocument(
            blocks=kept,
            title=_title_of(document, blocks, hint),
            doc_class=DocClass.OFFICE_DOC,
            parser=self.spec.name,
            parser_version=self.spec.version,
            truncated=truncated,
            meta={"paragraphs": index},
        )


def _paragraph_block(paragraph: Any, stack: HeadingStack, index: int) -> Block | None:
    text = clean_text(paragraph.text)

    if not text:
        return None

    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    heading = _HEADING_STYLE.match(style.strip())

    if heading:
        level = int(heading.group(1))
        stack.push(level, text)

        return Block(
            text=text,
            kind=BlockKind.HEADING,
            level=level,
            path=stack.current(),
            locator=f"¶{index}",
        )

    kind = BlockKind.LIST_ITEM if _LIST_STYLE.search(style) else BlockKind.PARAGRAPH

    return Block(
        text=text,
        kind=kind,
        path=stack.current(),
        locator=f"¶{index}",
    )


def _table_block(table: Any, stack: HeadingStack, index: int) -> Block | None:
    """把一张表渲染成 markdown 表格。

    为什么是 markdown 而不是"每行一句自然语言": 表格的价值在于
    行列对齐关系, markdown 用最少的字符把它保住了, 而且 LLM 对这种
    形态的理解相当好。逐行展开成 "姓名是张三, 年龄是 30" 则会把
    token 数翻几倍, 还是在重复表头。
    """

    rows: list[list[str]] = []

    for row in table.rows:
        cells = [clean_text(cell.text).replace("\n", " ").replace("|", "\\|")
                 for cell in row.cells]

        if any(cells):
            rows.append(cells)

    if not rows:
        return None

    width = max(len(r) for r in rows)
    lines = ["| " + " | ".join(_pad(rows[0], width)) + " |"]
    lines.append("| " + " | ".join(["---"] * width) + " |")

    for row in rows[1:]:
        lines.append("| " + " | ".join(_pad(row, width)) + " |")

    return Block(
        text="\n".join(lines),
        kind=BlockKind.TABLE,
        path=stack.current(),
        locator=f"表 @¶{index}",
        meta={"rows": len(rows), "cols": width},
    )


def _pad(row: list[str], width: int) -> list[str]:
    return row + [""] * (width - len(row))


def _title_of(document: Any, blocks: list[Block], hint: ParseHint) -> str:
    try:
        core = document.core_properties
        title = clean_text(core.title or "")
    except Exception:
        title = ""

    if title:
        return title

    first_heading = next((b.text for b in blocks if b.is_heading), None)

    return first_heading or hint.filename


PARSER = WordParser()

__all__ = ["PARSER", "WordParser"]
