"""六种文档解析器的测试。

样本文件全部在测试里现场生成, 不放二进制夹具: 夹具会随着库版本升级
慢慢变成"没人知道它当初为什么长这样"的黑盒, 而且 review 一个 .docx
的 diff 等于没有 review。

每个解析器只验证两件事: 文本有没有被完整取出来, 以及**结构有没有被
保住** —— 后者是这一层存在的理由。如果只要文本, 下游就只剩按字数硬切
一条路, 章节、页码、单元格位置全部丢失。
"""

from __future__ import annotations

import io

import pytest

from shovel.domain.enums import BlockKind, DocClass
from shovel.exceptions.pipeline import (
    EmptyDocumentError,
    ParseError,
    UnsupportedDocumentError,
)
from shovel.pipeline.parsers import ParseHint, can_parse, parse_bytes, supported_suffixes


def _hint(filename: str) -> ParseHint:
    return ParseHint(uri=f"file:///t/{filename}", filename=filename)


def _parse(filename: str, data: bytes | str):
    if isinstance(data, str):
        data = data.encode("utf-8")

    return parse_bytes(data, _hint(filename))


def _text(parsed) -> str:
    return "\n".join(b.text for b in parsed.blocks)


# --------------------------------------------------------------------------- #
# 路由
# --------------------------------------------------------------------------- #
def test_registry_covers_the_six_required_formats() -> None:
    for suffix in (".txt", ".md", ".pdf", ".docx", ".xlsx", ".pptx"):
        assert suffix in supported_suffixes()
        assert can_parse(_hint(f"a{suffix}"))


def test_unknown_suffix_is_refused_explicitly() -> None:
    """拒绝必须是一个明确的异常, 而不是"解析出空文档"。

    后者会让 .exe / .zip 静悄悄进到 document 表里, 用户永远查不出为什么
    索引里有一堆乱码。
    """

    with pytest.raises(UnsupportedDocumentError):
        _parse("setup.exe", b"MZ\x90\x00binary")


def test_legacy_office_gets_an_actionable_hint() -> None:
    """.doc/.xls/.ppt 是完全不同的二进制格式, 现代库读不了。

    重点是报错里要带"另存为"的建议 —— 用户能自己解决, 就不该让他
    去翻文档。
    """

    with pytest.raises(UnsupportedDocumentError) as err:
        _parse("旧报告.doc", b"\xd0\xcf\x11\xe0")

    assert err.value.hint


# --------------------------------------------------------------------------- #
# 纯文本
# --------------------------------------------------------------------------- #
def test_plain_text_splits_on_blank_lines() -> None:
    parsed = _parse("a.txt", "第一段。\n继续第一段。\n\n第二段。\n")

    assert parsed.doc_class is DocClass.PLAIN_TEXT
    assert len(parsed.blocks) == 2
    assert "继续第一段" in parsed.blocks[0].text


def test_plain_text_decodes_gb18030() -> None:
    """国内用户的历史文件大量是 GBK。猜成 latin-1 会得到一堆问号,

    而那种损坏不会报错 —— 它会一路写进索引, 直到用户搜不到为止。
    """

    parsed = _parse("gbk.txt", "中文编码测试内容。".encode("gb18030"))

    assert "中文编码测试内容" in _text(parsed)


def test_binary_disguised_as_text_is_refused() -> None:
    """NUL 字节是"这不是文本"的可靠信号。"""

    with pytest.raises((ParseError, UnsupportedDocumentError)):
        _parse("fake.txt", b"\x00\x01\x02\x00\x03binary")


def test_empty_file_is_refused() -> None:
    with pytest.raises(EmptyDocumentError):
        _parse("blank.txt", "   \n\n  \n")


# --------------------------------------------------------------------------- #
# Markdown
# --------------------------------------------------------------------------- #
def test_markdown_keeps_heading_hierarchy() -> None:
    """标题层级是 context_prefix 的唯一来源, 丢了它检索块就没有出处。"""

    parsed = _parse(
        "a.md",
        "# 一级标题\n\n正文甲。\n\n## 二级标题\n\n正文乙。\n",
    )

    headings = [b for b in parsed.blocks if b.kind is BlockKind.HEADING]
    assert [b.level for b in headings] == [1, 2]

    body = [b for b in parsed.blocks if b.kind is BlockKind.PARAGRAPH][-1]
    assert body.path == ("一级标题", "二级标题")


def test_markdown_code_fence_is_one_atomic_block() -> None:
    """代码块从中间切开就不再是可运行的代码, 检索出来也没用。"""

    parsed = _parse(
        "a.md",
        "说明文字。\n\n```python\ndef f():\n    return 1\n\n\ndef g():\n    return 2\n```\n",
    )

    code = [b for b in parsed.blocks if b.kind is BlockKind.CODE]
    assert len(code) == 1
    assert "def g()" in code[0].text
    assert code[0].kind.atomic


def test_markdown_front_matter_does_not_become_body() -> None:
    parsed = _parse("a.md", "---\ntitle: 笔记\ntags: [a, b]\n---\n\n真正的正文。\n")

    assert "tags: [a, b]" not in _text(parsed)
    assert "真正的正文" in _text(parsed)


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _pdf(text: str) -> bytes:
    """手工拼一个最小合法 PDF。

    用 reportlab 生成会多一个只为测试而存在的依赖; 这几十行虽然丑,
    但它精确地只包含被测代码需要的东西。
    """

    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []

    for i, obj in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"

    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()

    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()

    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()

    return bytes(out)


def test_pdf_extracts_text_with_page_locator() -> None:
    """页码是 PDF 唯一有意义的定位单位, 引用回原文全靠它。"""

    parsed = _parse("a.pdf", _pdf("Hello from page one"))

    assert parsed.doc_class is DocClass.PDF_DOC
    assert "Hello from page one" in _text(parsed)
    assert parsed.blocks[0].locator == "p.1"


def test_scanned_pdf_is_reported_as_empty_not_success() -> None:
    """扫描件有页面却没有文本层。当成"成功解析出空文档"的话,

    用户会看到一篇 indexed 的文档却一个字都搜不到, 而且没有任何线索。
    """

    with pytest.raises(EmptyDocumentError):
        _parse("scan.pdf", _pdf(""))


def test_corrupt_pdf_raises_parse_error() -> None:
    with pytest.raises(ParseError):
        _parse("broken.pdf", b"definitely not a pdf at all")


# --------------------------------------------------------------------------- #
# Word
# --------------------------------------------------------------------------- #
def _docx() -> bytes:
    import docx

    buf = io.BytesIO()
    doc = docx.Document()
    doc.add_heading("第一章 总则", 1)
    doc.add_paragraph("正文段落甲。")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "服务费"
    table.cell(1, 1).text = "1000"

    doc.add_paragraph("表格之后的正文段落乙。")
    doc.save(buf)

    return buf.getvalue()


def test_docx_preserves_body_order_across_tables() -> None:
    """python-docx 的 paragraphs 和 tables 是两个独立列表。

    照着它们直接遍历, 所有表格都会被挪到文末 —— 合同里"下表所列"
    的那段话和表格会相隔几十段, 语义彻底断掉。
    """

    parsed = _parse("a.docx", _docx())
    texts = [b.text for b in parsed.blocks]

    table_at = next(i for i, b in enumerate(parsed.blocks) if b.kind is BlockKind.TABLE)
    before = next(i for i, t in enumerate(texts) if "正文段落甲" in t)
    after = next(i for i, t in enumerate(texts) if "正文段落乙" in t)

    assert before < table_at < after


def test_docx_heading_level_comes_from_style() -> None:
    parsed = _parse("a.docx", _docx())
    heading = parsed.blocks[0]

    assert heading.kind is BlockKind.HEADING
    assert heading.level == 1
    assert "第一章" in heading.text


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def _xlsx(*, hidden: bool = False) -> bytes:
    import openpyxl

    buf = io.BytesIO()
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "订单"
    sheet.append(["商品", "数量", "单价"])

    for i in range(1, 6):
        sheet.append([f"商品{i}", i, i * 10])

    if hidden:
        backup = book.create_sheet("旧备份")
        backup.append(["早就删掉的数据"])
        backup.sheet_state = "hidden"

    book.save(buf)

    return buf.getvalue()


def test_xlsx_repeats_header_in_every_block() -> None:
    """表格切块后, 不带表头的那一半就是一串无意义的数字。"""

    parsed = _parse("a.xlsx", _xlsx())

    assert parsed.doc_class is DocClass.TABLE
    assert all("商品" in b.text for b in parsed.blocks)
    assert "订单" in (parsed.blocks[0].locator or "")


def test_xlsx_skips_hidden_sheets() -> None:
    """隐藏表通常是中间过程或历史备份, 索引它等于让用户搜到已删数据。"""

    parsed = _parse("a.xlsx", _xlsx(hidden=True))

    assert "早就删掉的数据" not in _text(parsed)
    assert "旧备份" in parsed.meta["sheets"]


def test_xlsx_meta_survives_workbook_close() -> None:
    """read_only 模式下 close() 会释放 zip 句柄。

    如果 meta 是在关闭之后才去读 book.worksheets, 每一个 xlsx 都会炸。
    """

    parsed = _parse("a.xlsx", _xlsx())

    assert parsed.meta["sheets"] == ["订单"]


# --------------------------------------------------------------------------- #
# PowerPoint
# --------------------------------------------------------------------------- #
def _pptx() -> bytes:
    from pptx import Presentation

    buf = io.BytesIO()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "季度总结"
    slide.placeholders[1].text = "要点一\n要点二"
    slide.notes_slide.notes_text_frame.text = "这里是演讲者备注。"
    deck.save(buf)

    return buf.getvalue()


def test_pptx_keeps_one_block_per_slide() -> None:
    parsed = _parse("a.pptx", _pptx())

    assert "季度总结" in _text(parsed)
    assert "要点二" in _text(parsed)
    assert any((b.locator or "").startswith("slide") for b in parsed.blocks)


def test_pptx_captures_speaker_notes() -> None:
    """备注里往往才是真正的论证过程, 幻灯片上只有关键词。"""

    parsed = _parse("a.pptx", _pptx())
    notes = [b for b in parsed.blocks if b.kind is BlockKind.NOTE]

    assert notes
    assert "演讲者备注" in notes[0].text
